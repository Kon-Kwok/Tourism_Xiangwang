#!/usr/bin/env python3
"""从导出的 Excel 中提取「店铺关键数据完成情况」sheet 第1-19行，生成 HTML 邮件正文。

用法:
    python3 scripts/generate_email_body.py exports/象往日报.xlsx         # 输出 HTML 到 stdout
    python3 scripts/generate_email_body.py exports/象往日报.xlsx --send   # 直接发送邮件

SMTP 配置通过环境变量:
    SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASS, MAIL_FROM, MAIL_TO
"""
from __future__ import annotations

import os
import smtplib
import sys
from email.mime.text import MIMEText
from pathlib import Path

import openpyxl

FOOTER = (
    "请查收皇家加勒比飞猪旗舰店全店数据（店铺基础数据、赤兔统计数据、阿里妈妈推广等）。"
    "本次日报数据均采用AI数据采集工具批量自动化生成和发送，如有偏差，"
    "请随时联系项目经理或邮箱：rcclapac@the-shineon.com，我们将第一时间修正系统，谢谢。"
)

DARK_BLUE = "#002060"
GREEN = "#008000"
RED = "#FF0000"
BORDER_COLOR = "#cccccc"
BG_WHITE = "#ffffff"
BG_STRIPE = "#f5f7fa"


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        if isinstance(value, float) and -1 < value < 1:
            return f"{value:+.2%}"
        if value == int(value):
            return f"{int(value):,}"
        return f"{value:,.2f}"
    return str(value)


def resolve_formula(ws, ref_str: str):
    try:
        if ref_str.startswith("="):
            ref_str = ref_str[1:]
        if ref_str.startswith("IFERROR("):
            inner = ref_str[len("IFERROR("):].rsplit(",", 1)[0]
            if "/" in inner:
                a, b = inner.split("/", 1)
                av = resolve_formula(ws, a)
                bv = resolve_formula(ws, b)
                if av is not None and bv:
                    return av / bv
            return None
        if ref_str.startswith("SUM("):
            inner = ref_str[len("SUM("):].rstrip(")")
            sr, er = inner.split(":")
            sc, sr_num = sr[0], int(sr[1:])
            ec, er_num = er[0], int(er[1:])
            total = 0
            for rr in range(sr_num, er_num + 1):
                for cc in range(ord(sc) - 64, ord(ec) - 64 + 1):
                    v = ws.cell(row=rr, column=cc).value
                    if isinstance(v, (int, float)):
                        total += v
            return total
        c = ws[ref_str]
        v = c.value
        if v is None:
            return None
        if isinstance(v, (int, float)):
            return v
        if isinstance(v, str) and v.startswith("="):
            return resolve_formula(ws, v[1:])
        return None
    except Exception:
        return None


def _resolve_cell(ws, row: int, col: int) -> tuple[str, str | None]:
    """Return (display_text, color) for a cell."""
    cell = ws.cell(row=row, column=col)
    v = cell.value
    if v is None:
        return "", None

    s = str(v)
    color = None

    if s.startswith("=IFERROR") or s.startswith("=SUM") or (s.startswith("=") and row >= 13):
        resolved = resolve_formula(ws, s)
        if resolved is not None:
            is_rate = s.startswith("=IFERROR") or "完成率" in str(ws.cell(row=row, column=1).value or "")
            if is_rate:
                text = f"{resolved:.1%}"
            elif isinstance(resolved, float) and -1 < resolved < 1:
                text = f"{resolved:+.2%}"
            elif isinstance(resolved, float):
                text = f"{resolved:,.2f}" if resolved != int(resolved) else f"{int(resolved):,}"
            else:
                text = str(resolved)
            if not is_rate and isinstance(resolved, float) and -1 < resolved < 1 and resolved != 0:
                color = GREEN if resolved > 0 else RED
            return text, color
        else:
            return cell_text(v), None

    if isinstance(v, float) and -1 < v < 1:
        text = f"{v:+.2%}"
        color = GREEN if v > 0 else RED if v < 0 else None
        return text, color

    return cell_text(v), None


# ---- HTML helpers ----

CSS = """<style>
  body { margin:0; padding:0; font-family:'Microsoft YaHei','PingFang SC',Arial,sans-serif; font-size:14px; color:#333; }
  table { border-collapse:collapse; }
</style>"""


def _th(text: str, colspan: int = 1) -> str:
    return (
        f'<th colspan="{colspan}" style="'
        f'background:{DARK_BLUE}; color:#fff; font-weight:bold; '
        f'padding:8px 12px; border:1px solid {DARK_BLUE}; text-align:center; '
        f'font-size:13px; white-space:nowrap;">{text}</th>'
    )


def _td(text: str, *, color: str | None = None, bold: bool = False,
        align: str = "center", bg: str | None = None, nowrap: bool = False) -> str:
    styles = [
        f"padding:6px 10px",
        f"border:1px solid {BORDER_COLOR}",
        f"text-align:{align}",
        f"font-size:13px",
        f"background:{bg or BG_WHITE}",
    ]
    if bold:
        styles.append("font-weight:bold")
    if color:
        styles.append(f"color:{color}")
    if nowrap:
        styles.append("white-space:nowrap")
    return f'<td style="{"; ".join(styles)}">{text}</td>'


# ---- Main HTML generation ----

def generate_daily_report_html(ws) -> str:
    """Generate complete HTML string for the email body."""
    date_val = cell_text(ws.cell(row=1, column=2).value)
    parts = []

    parts.append(f"""<!DOCTYPE html>
<html lang="zh-CN">
<head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1.0">
<title>店铺关键数据完成情况 — {date_val}</title>
{CSS}
</head>
<body style="max-width:700px; margin:0 auto; padding:8px;">""")

    parts.append(
        f'<p style="margin:0 0 10px 0; font-size:15px; font-weight:bold;">'
        f'店铺关键数据完成情况 &nbsp; Date: {date_val}</p>'
    )

    # === 店铺数据 (Rows 4-6) ===
    parts.append('<table style="width:100%; margin-bottom:12px;">')
    hdr = "".join(_th(cell_text(ws.cell(row=4, column=c).value)) for c in range(1, 8))
    parts.append(f"<tr>{hdr}</tr>")
    for r, bg in ((5, BG_WHITE), (6, BG_STRIPE)):
        cells = []
        for c in range(1, 8):
            text, color = _resolve_cell(ws, r, c)
            cells.append(_td(text, color=color, bold=(c == 1), align="left" if c == 1 else "center", bg=bg, nowrap=(c > 1)))
        parts.append(f"<tr>{''.join(cells)}</tr>")
    parts.append("</table>")

    # === 客服数据 (Rows 8-10) ===
    parts.append('<table style="width:100%; margin-bottom:12px;">')
    hdr = "".join(_th(cell_text(ws.cell(row=8, column=c).value)) for c in range(1, 7))
    parts.append(f"<tr>{hdr}</tr>")
    for r, bg in ((9, BG_WHITE), (10, BG_STRIPE)):
        cells = []
        for c in range(1, 7):
            text, color = _resolve_cell(ws, r, c)
            cells.append(_td(text, color=color, bold=(c == 1), align="left" if c == 1 else "center", bg=bg, nowrap=(c > 1)))
        parts.append(f"<tr>{''.join(cells)}</tr>")
    parts.append("</table>")

    # === MTD / YTD (Rows 12-19) ===
    parts.append('<table style="width:100%; margin-bottom:12px;">')
    parts.append(f"<tr>{_th('Month Target', colspan=2)}</tr>")
    for r in range(13, 20):
        label_text, _ = _resolve_cell(ws, r, 1)
        val_text, val_color = _resolve_cell(ws, r, 2)
        if not label_text:
            continue
        bg = BG_STRIPE if (r % 2 == 0) else BG_WHITE
        parts.append(
            f"<tr>"
            f'{_td(label_text, bold=True, align="left", bg=bg)}'
            f'{_td(val_text, color=val_color, bg=bg)}'
            f"</tr>"
        )
    parts.append("</table>")

    # Footer
    parts.append(
        f'<p style="color:#888; font-size:12px; margin-top:16px; line-height:1.6;">{FOOTER}</p>'
    )
    parts.append("</body></html>")

    return "\n".join(parts)


# ---- Mail sending ----

def send_mail(html_body: str, subject: str) -> bool:
    """Send HTML email via SMTP."""
    smtp_host = os.environ.get("SMTP_HOST", "")
    smtp_port = int(os.environ.get("SMTP_PORT", "587"))
    smtp_user = os.environ.get("SMTP_USER", "")
    smtp_pass = os.environ.get("SMTP_PASS", "")
    mail_from = os.environ.get("MAIL_FROM", smtp_user)
    mail_to = os.environ.get("MAIL_TO", "").split(",")

    if not all([smtp_host, smtp_user, smtp_pass, mail_to[0]]):
        print("错误: 缺少 SMTP 环境变量 (SMTP_HOST, SMTP_USER, SMTP_PASS, MAIL_TO)", file=sys.stderr)
        return False

    msg = MIMEText(html_body, "html", "utf-8")
    msg["Subject"] = subject
    msg["From"] = mail_from
    msg["To"] = ", ".join(mail_to)

    try:
        with smtplib.SMTP(smtp_host, smtp_port, timeout=30) as server:
            server.starttls()
            server.login(smtp_user, smtp_pass)
            server.sendmail(mail_from, mail_to, msg.as_string())
        print(f"邮件已发送: {subject} -> {mail_to}", file=sys.stderr)
        return True
    except Exception as e:
        print(f"发送失败: {e}", file=sys.stderr)
        return False


# ---- CLI ----

def main():
    flags = [a for a in sys.argv[1:] if a.startswith("--")]
    args = [a for a in sys.argv[1:] if not a.startswith("--")]

    if len(args) < 1:
        print("用法: python3 scripts/generate_email_body.py <excel文件> [--send]", file=sys.stderr)
        sys.exit(1)

    xlsx_path = Path(args[0])
    if not xlsx_path.exists():
        print(f"错误: 文件不存在 - {xlsx_path}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(xlsx_path)
    if "店铺关键数据完成情况" not in wb.sheetnames:
        print("错误: 未找到「店铺关键数据完成情况」sheet", file=sys.stderr)
        sys.exit(1)

    ws = wb["店铺关键数据完成情况"]
    date_val = cell_text(ws.cell(row=1, column=2).value)
    subject = f"象旺日报 — 店铺关键数据完成情况 ({date_val})"

    html_body = generate_daily_report_html(ws)

    if "--send" in flags:
        success = send_mail(html_body, subject)
        sys.exit(0 if success else 1)
    else:
        print(html_body)


if __name__ == "__main__":
    main()
